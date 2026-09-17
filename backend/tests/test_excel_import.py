import tempfile
import math
from pathlib import Path
import unittest

import openpyxl

from apps.datasource.utils.excel_import import (
    ConversionStore, ImportValidationError, convert_sheet, import_conversion, typed_frame,
)


class ExcelImportTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'input.xlsx'
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = '综测'
        sheet.append(['学号', '基本素质评价B', None])
        sheet.append([None, '思想品德B1 / 政治素养B11', None])
        sheet.append([None, 'b11分值', '加减分项目'])
        sheet.append(['001', 0, '含\t制表符\n换行'])
        sheet.append(['002', 4.5, "'备注"])
        sheet.append([None, None, None])
        sheet.append(['003', None, '=literal'])
        sheet['C7'].data_type = 's'
        sheet.merge_cells('A1:A3')
        sheet.merge_cells('B1:C1')
        sheet.merge_cells('B2:C2')
        normal = book.create_sheet('普通')
        normal.append(['姓名', '数值'])
        normal.append(['学生', 2])
        book.save(self.source)
        self.option = {'sheetName': '综测', 'headerMode': 'multi', 'headerStartRow': 1,
                       'headerEndRow': 3, 'dataStartRow': 4}

    def test_hierarchy_values_and_types(self):
        sheet = convert_sheet(self.source, self.option)
        self.assertEqual([f['fieldName'] for f in sheet['fields']], ['学号', 'B11_score', 'B11_items'])
        self.assertIn('基本素质评价B', sheet['fields'][1]['meaning'])
        self.assertEqual(sheet['skippedEmptyRows'], [6])
        self.assertEqual(sheet['sourceRows'], [4, 5, 7])
        frame = typed_frame(sheet, {})
        self.assertEqual(frame.iloc[0, 0], '001')
        self.assertEqual(frame.iloc[0, 1], 0)
        self.assertEqual(frame.iloc[2, 2], '=literal')
        self.assertEqual(len(frame), 3)

    def test_mixed_sheets_preview_and_owned_mapping(self):
        store = ConversionStore(self.root, 1, 1)
        identifier, _ = store.upload('input.xlsx', self.source.read_bytes())
        options = [self.option, {'sheetName': '普通'}]
        conversion_id, result = store.convert(identifier, options)
        self.assertEqual(store.convert(identifier, options)[0], conversion_id)
        preview = store.preview(conversion_id, result)
        self.assertEqual(preview['data'][0]['data'][0]['学号'], '001')
        self.assertEqual(preview['data'][1]['fields'][0]['fieldName'], '姓名')
        self.assertEqual(store.load(conversion_id, 'conversion')['sheets'], result['sheets'])
        for other in (ConversionStore(self.root, 2, 1), ConversionStore(self.root, 1, 2)):
            with self.assertRaises(ImportValidationError):
                other.load(conversion_id, 'conversion')
        with self.assertRaises(ImportValidationError):
            store.load('../input', 'upload')

    def test_upload_ignores_empty_sheets(self):
        book = openpyxl.load_workbook(self.source)
        book.create_sheet('空白模板')
        book.save(self.source)
        store = ConversionStore(self.root, 1, 1)
        upload_id, upload = store.upload('input.xlsx', self.source.read_bytes())
        self.assertEqual(upload['sheetNames'], ['综测', '普通'])
        conversion_id, conversion = store.convert(upload_id, [
            {'sheetName': '普通', 'headerMode': 'single'}
        ])
        self.assertEqual([sheet['sheetName'] for sheet in conversion['sheets']], ['普通'])
        self.assertTrue(conversion_id)

    def test_upload_rejects_workbook_with_only_empty_sheets(self):
        book = openpyxl.Workbook()
        book.save(self.source)
        store = ConversionStore(self.root, 1, 1)
        with self.assertRaisesRegex(ImportValidationError, '没有可导入的数据'):
            store.upload('input.xlsx', self.source.read_bytes())

    def test_invalid_ranges_and_merged_data(self):
        for changes in ({'headerEndRow': 2}, {'dataStartRow': 2}, {'lastColumn': 2}, {'firstColumn': 9}):
            with self.subTest(changes=changes), self.assertRaises(ImportValidationError):
                convert_sheet(self.source, {**self.option, **changes})
        book = openpyxl.load_workbook(self.source)
        book['综测'].merge_cells('B4:B5')
        book.save(self.source)
        with self.assertRaisesRegex(ImportValidationError, '数据区'):
            convert_sheet(self.source, self.option)

    def test_missing_formula_cache_and_invalid_type(self):
        book = openpyxl.load_workbook(self.source)
        book['综测']['B4'] = '=1+2'
        book.save(self.source)
        with self.assertRaisesRegex(ImportValidationError, '缓存'):
            convert_sheet(self.source, self.option)
        sheet = convert_sheet(self.source, {'sheetName': '普通'})
        with self.assertRaises(ImportValidationError):
            typed_frame(sheet, {'col_1': 'int'})

    def test_duplicate_long_chinese_and_short_reserved_names(self):
        book = openpyxl.Workbook()
        sheet = book.active
        prefix = '能力素质评价A_社会工作A1 任职1*考核系数+任职2*考核系数*0.2 满分10分'
        sheet.append([prefix + '_A1', prefix + '_加减分项目', '姓名', '姓名'])
        sheet.append([1, '项目', '甲', '乙'])
        book.save(self.source)
        result = convert_sheet(self.source, {'sheetName': 'Sheet'})
        names = [f['fieldName'] for f in result['fields']]
        self.assertEqual(len(set(names)), 4)
        self.assertTrue(all(len(n.encode()) <= 63 for n in names))
        self.assertEqual(result['fields'][0]['meaning'], prefix + '_A1')
        self.assertTrue(result['warnings'])

    def test_csv_preserves_identifiers_and_rejects_multi(self):
        source = self.root / 'input.csv'
        source.write_text('学号,数值\n001,2\n002,0\n', encoding='utf-8')
        sheet = convert_sheet(source, {'sheetName': 'Sheet1'})
        self.assertEqual(typed_frame(sheet, {}).iloc[0, 0], '001')
        with self.assertRaises(ImportValidationError):
            convert_sheet(source, {**self.option, 'sheetName': 'Sheet1'})

    def test_forged_field_selection_rejected_before_connection(self):
        sheet = convert_sheet(self.source, self.option)
        with self.assertRaises(ImportValidationError):
            import_conversion(None, 'a' * 32, {'sheets': [sheet]}, [
                {'sheetName': '综测', 'fields': [{'fieldId': 'unknown', 'fieldType': 'int'}]}])

    def test_original_xls_matches_existing_flattened_data(self):
        root = Path(__file__).resolve().parents[2] / '项目学习文档' / '数据集和测试问题集'
        sources = list(root.glob('表1：*.xls'))
        if not sources:
            self.skipTest('Original local workbook is not part of the repository')
        from apps.datasource.utils.excel_import import read_sheet
        sheet = convert_sheet(sources[0], {'sheetName': 'Sheet2', 'headerMode': 'multi',
                              'headerStartRow': 1, 'headerEndRow': 4, 'dataStartRow': 5})
        expected, _, _ = read_sheet(next(root.glob('表1：*_导入版.xlsx')), '数据')
        self.assertEqual(len(sheet['values']), len(expected) - 1)
        for row_index, (actual, baseline) in enumerate(zip(sheet['values'], expected[1:]), start=5):
            self.assertEqual(len(actual), len(baseline))
            for col_index, (left, right) in enumerate(zip(actual, baseline), start=1):
                if isinstance(left, (int, float)) and isinstance(right, (int, float)):
                    matches = math.isclose(left, right, rel_tol=1e-12, abs_tol=1e-12)
                else:
                    matches = left == right or left in (None, '') and right in (None, '')
                self.assertTrue(matches, f'Value mismatch at source row {row_index}, column {col_index}')
        names = [field['fieldName'] for field in sheet['fields']]
        self.assertEqual(len(set(names)), len(names))
        self.assertTrue(all(len(name.encode('utf-8')) <= 63 for name in names))
        self.assertIn('A1_score', names)
        self.assertIn('A1_items', names)
        self.assertIn('B_score', names)


if __name__ == '__main__':
    unittest.main()
