"""User-directed Excel conversion. Stored mappings are authoritative at import time."""
import datetime as dt
import hashlib
import json
import math
import os
from pathlib import Path
import re
import uuid
from contextlib import contextmanager

import pandas as pd

from .identifiers import postgres_column_names


class ImportValidationError(ValueError):
    pass


def _encode(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return {'__excel_type': type(value).__name__, 'value': value.isoformat()}
    raise TypeError(f'Unsupported cell type: {type(value).__name__}')


def _decode(value):
    kinds = {'datetime': dt.datetime, 'date': dt.date, 'time': dt.time}
    if value.get('__excel_type') in kinds:
        return kinds[value['__excel_type']].fromisoformat(value['value'])
    return value


def _blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def _label(value):
    return '' if value is None else re.sub(r'\s+', ' ', str(value)).strip()


def sheet_names(source):
    source = Path(source)
    if source.suffix.lower() == '.csv':
        return ['Sheet1']
    if source.suffix.lower() == '.xls':
        import xlrd
        book = xlrd.open_workbook(source, on_demand=True)
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    import openpyxl
    book = openpyxl.load_workbook(source, read_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def nonempty_sheet_names(source):
    """List sheets containing values; empty template tabs are ignored."""
    names = []
    for name in sheet_names(source):
        rows, _, _ = read_sheet(source, name)
        if any(not _blank(value) for row in rows for value in row):
            names.append(name)
    return names


def read_sheet(source, name):
    """Return values, real merged ranges and invalid cells; coordinates are zero based."""
    source = Path(source)
    if source.suffix.lower() == '.csv':
        import csv
        with source.open(encoding='utf-8-sig', newline='') as stream:
            return list(csv.reader(stream)), [], {}
    if source.suffix.lower() == '.xls':
        import xlrd
        book = xlrd.open_workbook(source, formatting_info=True)
        try:
            sheet = book.sheet_by_name(name)
            rows, errors = [], {}
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    value = cell.value
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        errors[r, c] = 'Excel 错误值'
                    row.append(value)
                rows.append(row)
            return rows, list(sheet.merged_cells), errors
        finally:
            book.release_resources()
    import openpyxl
    raw = openpyxl.load_workbook(source, data_only=False)
    cached = openpyxl.load_workbook(source, data_only=True)
    try:
        sheet, values = raw[name], cached[name]
        rows, errors = [], {}
        for row in sheet:
            result = []
            for cell in row:
                value = values.cell(cell.row, cell.column).value
                if cell.data_type == 'f' and value is None:
                    errors[cell.row - 1, cell.column - 1] = '公式缺少缓存结果，请在 Excel/WPS 重算并保存'
                elif cell.data_type == 'e' or values.cell(cell.row, cell.column).data_type == 'e':
                    errors[cell.row - 1, cell.column - 1] = 'Excel 错误值'
                result.append(value)
            rows.append(result)
        merges = [(m.min_row - 1, m.max_row, m.min_col - 1, m.max_col)
                  for m in sheet.merged_cells.ranges]
        return rows, merges, errors
    finally:
        raw.close()
        cached.close()


def short_name(levels):
    leaf = levels[-1]
    code = None
    for label in levels:
        # Formulas describe the indicator; their operands are not its code.
        title = re.split(r'[=（(]', label, maxsplit=1)[0]
        matches = re.findall(r'(?<![A-Za-z0-9])([A-Za-z]\d{0,3})(?![A-Za-z0-9])', title)
        if matches:
            code = matches[-1].upper()
    if '加减分' in leaf or '减分项目' in leaf:
        kind = 'items'
    elif '分值' in leaf or '总分' in leaf or re.fullmatch(r'[A-Za-z]\d{0,3}', leaf):
        kind = 'score'
    else:
        kind = re.sub(r'[^\w\u4e00-\u9fff]', '_', leaf).strip('_')
    if code:
        return f'{code}_{kind or "value"}'
    return '_'.join(levels[-2:])


def infer_type(values):
    nonempty = [v for v in values if not _blank(v)]
    if not nonempty:
        return 'string'
    if all(isinstance(v, (dt.datetime, dt.date)) for v in nonempty):
        return 'datetime'
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in nonempty):
        return 'int' if all(float(v).is_integer() for v in nonempty) else 'float'
    # CSV numeric values may be inferred, but digit strings with leading zeros stay text.
    if all(isinstance(v, str) for v in nonempty):
        if any(re.fullmatch(r'[+-]?0\d+', v.strip()) for v in nonempty):
            return 'string'
        if all(re.fullmatch(r'[+-]?\d+', v.strip()) for v in nonempty):
            return 'int'
        if all(re.fullmatch(r'[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?', v.strip()) for v in nonempty):
            return 'float'
    return 'string'


def convert_sheet(source, option):
    name = option['sheetName']
    multi = option.get('headerMode', 'single') == 'multi'
    if multi and Path(source).suffix.lower() == '.csv':
        raise ImportValidationError('CSV 仅支持普通表头模式')
    rows, merges, errors = read_sheet(source, name)
    # Ignore trailing formatted but empty cells when determining the effective range.
    nonempty = [(r, c) for r, row in enumerate(rows) for c, v in enumerate(row) if not _blank(v)]
    if not nonempty:
        raise ImportValidationError(f'{name}: 工作表为空')
    height = max(r for r, _ in nonempty) + 1
    width = max(c for _, c in nonempty) + 1
    start = option.get('headerStartRow') if multi else 1
    end = option.get('headerEndRow') if multi else 1
    data_start = option.get('dataStartRow') if multi else 2
    stop = option.get('dataEndRow') or height
    first = option.get('firstColumn') or 1
    last = option.get('lastColumn') or width
    if not all(isinstance(v, int) and not isinstance(v, bool) for v in (start, end, data_start, stop, first, last)):
        raise ImportValidationError(f'{name}: 请填写表头起止行和数据起始行')
    if not (1 <= start <= end < data_start <= stop <= height and 1 <= first <= last <= width):
        raise ImportValidationError(f'{name}: 表头、数据或列范围无效')
    for (r, c), message in errors.items():
        if first - 1 <= c < last and (start - 1 <= r < end or data_start - 1 <= r < stop):
            raise ImportValidationError(f'{name}: 第{r + 1}行第{c + 1}列 {message}')
    matrix = [list(row) + [None] * (width - len(row)) for row in rows]
    for r0, r1, c0, c1 in merges:
        if c1 <= first - 1 or c0 >= last:
            continue
        if r0 < stop and r1 > data_start - 1:
            raise ImportValidationError(f'{name}: 数据区第{r0 + 1}行存在合并单元格')
        if r0 < end and r1 > start - 1:
            if r0 < start - 1 or r1 > end or c0 < first - 1 or c1 > last:
                raise ImportValidationError(f'{name}: 第{r0 + 1}行合并区域跨越所选表头边界')
            for r in range(r0, r1):
                for c in range(c0, c1):
                    matrix[r][c] = rows[r0][c0]
    paths = []
    for c in range(first - 1, last):
        levels = []
        for r in range(start - 1, end):
            label = _label(matrix[r][c])
            if label and (not levels or label != levels[-1]):
                levels.append(label)
        if not levels:
            raise ImportValidationError(f'{name}: 第{c + 1}列没有表头')
        if any('\x00' in label for label in levels):
            raise ImportValidationError(f'{name}: 第{c + 1}列表头包含 NUL 字符')
        paths.append(levels)
    names = postgres_column_names([short_name(p) if multi else p[0] for p in paths])
    data, source_rows, skipped = [], [], []
    for r in range(data_start - 1, stop):
        values = matrix[r][first - 1:last]
        if all(_blank(v) for v in values):
            skipped.append(r + 1)
            continue
        data.append(values)
        source_rows.append(r + 1)
    fields = [{'fieldId': f'col_{c + first}', 'sourceColumnIndex': c + first,
               'originalHeaderPath': levels, 'fieldName': names[c],
               'meaning': ' / '.join(levels), 'fieldType': infer_type([row[c] for row in data])}
              for c, levels in enumerate(paths)]
    warnings = []
    if len({tuple(p) for p in paths}) != len(paths):
        warnings.append('存在相同的完整表头，已分配不同列名；请确认这些列的业务含义。')
    return {'sheetName': name, 'fields': fields, 'values': data, 'sourceRows': source_rows,
            'skippedEmptyRows': skipped, 'rows': len(data), 'warnings': warnings,
            'range': {'headerStartRow': start, 'headerEndRow': end, 'dataStartRow': data_start,
                      'dataEndRow': stop, 'firstColumn': first, 'lastColumn': last}}


def typed_frame(sheet, types):
    result = {}
    for index, field in enumerate(sheet['fields']):
        kind = types.get(field['fieldId'], field['fieldType'])
        values = [row[index] for row in sheet['values']]
        try:
            if kind == 'string':
                result[field['fieldName']] = pd.Series(values, dtype='string')
            elif kind in ('int', 'float'):
                values = [None if _blank(v) else v for v in values]
                numbers = pd.to_numeric(pd.Series(values), errors='raise')
                if any(not math.isfinite(float(v)) for v in numbers.dropna()):
                    raise ValueError('非有限数值')
                result[field['fieldName']] = numbers.astype('Int64' if kind == 'int' else 'Float64')
            elif kind == 'datetime':
                if any(isinstance(v, (int, float)) for v in values if v is not None):
                    raise ValueError('数字不能直接作为日期，请在原文件指定日期类型')
                result[field['fieldName']] = pd.to_datetime(
                    pd.Series([None if _blank(v) else v for v in values]), errors='raise')
            else:
                raise ValueError('不支持的字段类型')
        except (ValueError, TypeError, OverflowError) as error:
            raise ImportValidationError(f"{sheet['sheetName']}: 字段 {field['fieldName']} 无法转换为 {kind}: {error}") from error
    return pd.DataFrame(result)


class ConversionStore:
    def __init__(self, root, user_id, workspace_id):
        self.root = Path(root) / '_conversions'
        self.root.mkdir(parents=True, exist_ok=True)
        self.owner = [str(user_id), str(workspace_id or 1)]

    def _path(self, identifier):
        if not re.fullmatch(r'[a-f0-9]{32}', identifier or ''):
            raise ImportValidationError('文件或转换标识无效，请重新上传')
        return self.root / f'{identifier}.json'

    @contextmanager
    def lock(self, identifier):
        """Coordinate conversion/import with explicit cleanup across worker processes."""
        lock_path = self._path(identifier).with_suffix('.lock')
        with lock_path.open('a+b') as stream:
            if os.name == 'nt':
                import msvcrt
                stream.write(b'0')
                stream.flush()
                stream.seek(0)
                msvcrt.locking(stream.fileno(), msvcrt.LK_LOCK, 1)
            else:
                import fcntl
                fcntl.flock(stream, fcntl.LOCK_EX)
            try:
                yield
            finally:
                if os.name == 'nt':
                    stream.seek(0)
                    msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
                else:
                    fcntl.flock(stream, fcntl.LOCK_UN)

    def write(self, identifier, value):
        target = self._path(identifier)
        temporary = target.with_suffix(f'.{uuid.uuid4().hex}.tmp')
        try:
            temporary.write_text(json.dumps(value, ensure_ascii=False, default=_encode, allow_nan=False), encoding='utf-8')
            os.replace(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)

    def load(self, identifier, kind):
        try:
            value = json.loads(self._path(identifier).read_text(encoding='utf-8'), object_hook=_decode)
        except FileNotFoundError as error:
            raise ImportValidationError('文件或转换结果已失效，请重新上传') from error
        if value.get('owner') != self.owner or value.get('kind') != kind:
            raise ImportValidationError('无权访问此文件或转换结果')
        return value

    def upload(self, filename, contents):
        suffix = Path(filename).suffix.lower()
        if suffix not in ('.xlsx', '.xls', '.csv'):
            raise ImportValidationError('Only support .xlsx/.xls/.csv')
        if len(contents) > 50 * 1024 * 1024:
            raise ImportValidationError('文件不能超过 50 MB')
        identifier = uuid.uuid4().hex
        source = self.root / f'{identifier}{suffix}'
        source.write_bytes(contents)
        try:
            names = nonempty_sheet_names(source)
            if not names:
                raise ImportValidationError('文件中没有可导入的数据')
            metadata = {'kind': 'upload', 'owner': self.owner, 'filename': Path(filename).name,
                        'source': source.name, 'sha256': hashlib.sha256(contents).hexdigest(),
                        'sheetNames': names, 'createdAt': dt.datetime.now(dt.timezone.utc).isoformat()}
            self.write(identifier, metadata)
        except ImportValidationError:
            source.unlink(missing_ok=True)
            raise
        except Exception:
            source.unlink(missing_ok=True)
            raise ImportValidationError('文件无法解析，请确认是有效的 Excel/CSV 文件') from None
        return identifier, metadata

    def convert(self, upload_id, options):
        with self.lock(upload_id):
            return self._convert(upload_id, options)

    def _convert(self, upload_id, options):
        upload = self.load(upload_id, 'upload')
        selected = [o['sheetName'] for o in options]
        if not selected or len(selected) != len(set(selected)) or not set(selected) <= set(upload['sheetNames']):
            raise ImportValidationError('请选择有效且不重复的 Sheet')
        source = self.root / upload['source']
        if hashlib.sha256(source.read_bytes()).hexdigest() != upload['sha256']:
            raise ImportValidationError('原文件已改变，请重新上传')
        sheets = [convert_sheet(source, o) for o in options]
        for sheet in sheets:
            typed_frame(sheet, {})  # Fail before publishing a preview that cannot be imported.
        identifier = hashlib.sha256(json.dumps(
            [upload_id, options, 1], sort_keys=True, ensure_ascii=False
        ).encode()).hexdigest()[:32]
        value = {'kind': 'conversion', 'owner': self.owner, 'uploadId': upload_id,
                 'filename': upload['filename'], 'sourceSha256': upload['sha256'],
                 'ruleVersion': 1, 'options': options, 'sheets': sheets,
                 'createdAt': dt.datetime.now(dt.timezone.utc).isoformat()}
        # A separate, single-header data-value workbook is available for audit.
        import openpyxl
        book = openpyxl.Workbook()
        book.remove(book.active)
        for sheet in sheets:
            ws = book.create_sheet(sheet['sheetName'])
            for row in [[f['fieldName'] for f in sheet['fields']]] + sheet['values']:
                ws.append(row)
                for cell, item in zip(ws[ws.max_row], row):
                    if isinstance(item, str):
                        cell.data_type = 's'
            ws.freeze_panes = 'A2'
        target = self.root / f'{identifier}.xlsx'
        temporary = self.root / f'{identifier}.{uuid.uuid4().hex}.tmp.xlsx'
        try:
            book.save(temporary)
            os.replace(temporary, target)
        finally:
            book.close()
            temporary.unlink(missing_ok=True)
        value['copyFile'] = target.name
        self.write(identifier, value)
        return identifier, value

    def preview(self, identifier, value):
        data = []
        for sheet in value['sheets']:
            frame = typed_frame(sheet, {}).head(10)
            records = json.loads(frame.to_json(orient='records', date_format='iso'))
            data.append({k: v for k, v in sheet.items() if k not in ('values', 'sourceRows')})
            data[-1]['data'] = records
        return {'conversionId': identifier, 'filePath': value['uploadId'], 'data': data}


def cleanup_conversions(root, engine, retention_days=7, apply=False):
    """Explicit maintenance: only stale, never-imported files; dry-run by default."""
    from sqlalchemy import text
    if retention_days < 1:
        raise ImportValidationError('保留时间至少为 1 天')
    directory = Path(root) / '_conversions'
    if not directory.exists():
        return []
    cutoff = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=retention_days)
    removed = []
    # Upload lock -> conversion lock is the same order used by conversion.
    for upload_path in directory.glob('*.json'):
        try:
            upload = json.loads(upload_path.read_text(encoding='utf-8'))
        except FileNotFoundError:
            continue
        if upload.get('kind') != 'upload':
            continue
        store = ConversionStore(root, *upload['owner'])
        with store.lock(upload_path.stem):
            manifests = []
            for manifest_path in directory.glob('*.json'):
                manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
                if manifest.get('kind') == 'conversion' and manifest.get('uploadId') == upload_path.stem:
                    manifests.append((manifest_path, manifest))
            preserve_upload = False
            for manifest_path, manifest in manifests:
                with store.lock(manifest_path.stem):
                    created = dt.datetime.fromisoformat(manifest['createdAt'])
                    with engine.connect() as conn:
                        imported = conn.execute(text('''SELECT EXISTS (
                            SELECT 1 FROM pg_class c JOIN pg_namespace n ON n.oid=c.relnamespace
                            WHERE n.nspname=current_schema() AND obj_description(c.oid) LIKE :marker)'''),
                            {'marker': f'sqlbot-import:{manifest_path.stem}:%'}).scalar_one()
                    if imported or created >= cutoff:
                        preserve_upload = True
                        continue
                    removed.append(manifest_path.stem)
                    if apply:
                        (directory / f'{manifest_path.stem}.xlsx').unlink(missing_ok=True)
                        manifest_path.unlink()
            if not preserve_upload and dt.datetime.fromisoformat(upload['createdAt']) < cutoff:
                removed.append(upload_path.stem)
                if apply:
                    # Only generated basenames inside the managed directory may be removed.
                    source = directory / upload['source']
                    if source.parent.resolve() != directory.resolve() or not source.name.startswith(upload_path.stem):
                        raise ImportValidationError('原件路径不属于转换目录')
                    source.unlink(missing_ok=True)
                    upload_path.unlink()
    return removed


def import_conversion(engine, conversion_id, conversion, selections):
    """One transaction, one insert path. Advisory lock + table markers survive lost replies."""
    from sqlalchemy import text
    from psycopg2 import sql

    sheets = {s['sheetName']: s for s in conversion['sheets']}
    selected = [s['sheetName'] for s in selections]
    if not selected or len(set(selected)) != len(selected) or not set(selected) <= set(sheets):
        raise ImportValidationError('导入 Sheet 与转换结果不匹配')
    plans = []
    for selection in selections:
        sheet = sheets[selection['sheetName']]
        fields = selection['fields']
        by_name = {f['fieldName']: f['fieldId'] for f in sheet['fields']}
        types = {f.get('fieldId') or by_name.get(f.get('fieldName')): f['fieldType'] for f in fields}
        expected = {f['fieldId'] for f in sheet['fields']}
        if set(types) != expected or len(fields) != len(expected):
            raise ImportValidationError(f"{sheet['sheetName']}: 字段与转换结果不匹配，请重新预览")
        plans.append((sheet, types, typed_frame(sheet, types)))
    fingerprint = hashlib.sha256(json.dumps(
        [conversion_id, [(s['sheetName'], sorted(t.items())) for s, t, _ in plans]], ensure_ascii=False
    ).encode()).hexdigest()
    lock_id = int(fingerprint[:15], 16)
    results = []
    with engine.begin() as conn:
        conn.execute(text('SELECT pg_advisory_xact_lock(:key)'), {'key': lock_id})
        raw = conn.connection.driver_connection
        with raw.cursor() as cursor:
            for index, (sheet, _, frame) in enumerate(plans):
                token = hashlib.sha256(f'{fingerprint}:{index}'.encode()).hexdigest()[:20]
                prefix = ('excel_' + re.sub(r'[^\w]', '', sheet['sheetName'])).encode('utf-8')[:40].decode('utf-8', errors='ignore')
                table_name = f'{prefix}_{token}'
                marker = f'sqlbot-import:{conversion_id}:{fingerprint}:{index}'
                cursor.execute('SELECT obj_description(to_regclass(%s))', ('"' + table_name + '"',))
                existing = cursor.fetchone()[0]
                cursor.execute('SELECT to_regclass(%s)', ('"' + table_name + '"',))
                exists = cursor.fetchone()[0] is not None
                if exists:
                    if existing != marker:
                        raise ImportValidationError('导入目标表状态不匹配，请重新转换')
                else:
                    frame.to_sql(table_name, conn, if_exists='fail', index=False, chunksize=1000)
                    for field in sheet['fields']:
                        cursor.execute(sql.SQL('COMMENT ON COLUMN {}.{} IS %s').format(
                            sql.Identifier(table_name), sql.Identifier(field['fieldName'])), (field['meaning'],))
                    cursor.execute(sql.SQL('COMMENT ON TABLE {} IS %s').format(sql.Identifier(table_name)), (marker,))
                results.append({'sheetName': sheet['sheetName'], 'tableName': table_name,
                                'tableComment': sheet['sheetName'], 'rows': len(frame)})
    return {'filename': conversion['filename'], 'sheets': results, 'importId': fingerprint,
            'conversionId': conversion_id}
