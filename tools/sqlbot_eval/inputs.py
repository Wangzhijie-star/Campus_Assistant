from __future__ import annotations

import hashlib
import math
import re
from dataclasses import asdict
from pathlib import Path
from urllib.parse import urlparse

from .models import Case, Run, DIMENSIONS, now
from .storage import component, write_json, atomic_text, save_runs


def load_cases(path: Path, ids: list[str], sheet: str | None = None) -> list[Case]:
    import openpyxl
    if not ids or len(ids) != len(set(ids)):
        raise ValueError('必须显式指定不重复的题号')
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook[sheet] if sheet else workbook.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else '' for v in next(rows)]
        for required in ['序号', '原始问题', '预期答案']:
            if headers.count(required) != 1:
                raise ValueError(f'缺少或重复列：{required}')
        found = {}
        for rowno, row in enumerate(rows, 2):
            if all(v is None for v in row):
                continue
            values = {h: (str(v).strip() if v is not None else '') for h, v in zip(headers, row) if h}
            qid = values['序号']
            if not re.fullmatch(r'S\d{3,}', qid):
                raise ValueError(f'第 {rowno} 行题号应为 SXXX：{qid}')
            if qid in found:
                raise ValueError(f'重复题号：{qid}')
            if qid in ids and (not values['原始问题'] or not values['预期答案']):
                raise ValueError(f'{qid} 缺少问题或预期答案（也可能公式没有缓存）')
            found[qid] = Case(qid, values['原始问题'], values['预期答案'], values, rowno)
        missing = set(ids) - found.keys()
        if missing:
            raise ValueError(f'未找到题号：{sorted(missing)}')
        return [found[qid] for qid in ids]
    finally:
        workbook.close()


def validate_config(config: dict) -> dict:
    allowed = {'base_url', 'datasource_id', 'datasource_name', 'token_env', 'token_header',
               'timeout_s', 'connect_timeout_s', 'confirm_timeout_s', 'poll_s', 'dimensions',
               'weights', 'na_policy', 'max_prompt_chars', 'data_version', 'deployment_version',
               'model_version', 'rubric_version', 'sheet'}
    if set(config) - allowed:
        raise ValueError(f'不支持的配置字段：{sorted(set(config)-allowed)}；凭证请放环境变量')
    config = dict(config)
    url = urlparse(config.get('base_url', ''))
    if url.scheme not in {'http', 'https'} or not url.netloc or url.username or url.query or url.fragment:
        raise ValueError('base_url 必须为不带凭证和查询参数的 HTTP 地址，含实际 API 前缀')
    if type(config.get('datasource_id')) is not int or config['datasource_id'] <= 0:
        raise ValueError('datasource_id 必须是实际的正整数编号')
    if not config.get('datasource_name'):
        raise ValueError('必须配置数据源名称供核对')
    for key, default in [('timeout_s', 600), ('connect_timeout_s', 10), ('confirm_timeout_s', 30), ('poll_s', 2), ('max_prompt_chars', 150000)]:
        v = config.setdefault(key, default)
        if type(v) not in (int, float) or not math.isfinite(v) or v <= 0:
            raise ValueError(f'{key} 必须为有限正数')
    dimensions = config.setdefault('dimensions', DIMENSIONS.copy())
    if not dimensions or len(set(dimensions)) != len(dimensions) or any(d not in DIMENSIONS for d in dimensions):
        raise ValueError('评分维度必须是不重复的已支持维度代码')
    weights = config.setdefault('weights', {})
    if weights and (set(weights) != set(dimensions) or any(type(v) not in (int, float) or not math.isfinite(v) or v <= 0 for v in weights.values())):
        raise ValueError('权重必须覆盖所有维度，且为有限正数')
    if config.setdefault('na_policy', 'pending') not in {'pending', 'exclude'}:
        raise ValueError('na_policy 只能为 pending 或 exclude')
    config.setdefault('rubric_version', 'v1')
    return config


def prepare_batch(config: dict, questions: Path, ids: list[str], batch_dir: Path, rubric: Path, template: Path) -> Path:
    config = validate_config(config)
    cases = load_cases(questions, ids, config.get('sheet'))
    rules = rubric.read_text(encoding='utf-8-sig').strip()
    if not rules:
        raise ValueError('通用评分规则不能为空')
    template_bytes = template.read_bytes()
    import openpyxl
    with template.open('rb') as stream:
        book = openpyxl.load_workbook(stream, read_only=True)
        try:
            if not {'业务效果', '执行质量', '效率'} <= set(book.sheetnames):
                raise ValueError('成绩单模板缺少预期工作表')
        finally:
            book.close()
    component(batch_dir.name)
    batch_dir.mkdir(parents=True, exist_ok=False)
    manifest = {'schema_version': '1.0', 'batch_id': batch_dir.name, 'created_at': now(), 'config': config,
                'sources': {name: {'path': str(path.resolve()), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
                            for name, path in [('questions', questions), ('rubric', rubric), ('template', template)]}}
    write_json(batch_dir / 'manifest.json', manifest)
    write_json(batch_dir / 'cases.json', [asdict(c) for c in cases])
    atomic_text(batch_dir / 'rubric.txt', rules)
    atomic_text(batch_dir / 'prompt_template.txt', (Path(__file__).parent / 'templates/review_prompt.txt').read_text(encoding='utf-8'))
    (batch_dir / 'template.xlsx').write_bytes(template_bytes)
    save_runs(batch_dir, [Run(c.question_id) for c in cases])
    return batch_dir
